from __future__ import annotations

import hashlib
import io
import json
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Annotated

from fastapi import Depends, FastAPI, File, Header, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import delete, func, or_, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from .config import settings
from .database import SessionLocal, create_schema, get_session
from .models import (
    Asset,
    AuditEvent,
    BatchStatus,
    Finding,
    FindingEvent,
    FindingStatus,
    GovernanceSetting,
    ImportBatch,
    Observation,
    Remediation,
    RiskAcceptance,
    Role,
    Severity,
    Source,
    User,
    Verification,
    utcnow,
)
from .schemas import (
    APIImportRequest,
    AssetCreate,
    AssetPublic,
    AssetUpdate,
    AssignmentUpdate,
    AuditPublic,
    BatchPublic,
    EventPublic,
    FindingDetail,
    FindingSummary,
    FindingUpdate,
    GovernanceSettingPublic,
    GovernanceSettingUpdate,
    LoginRequest,
    ObservationInput,
    ObservationPublic,
    RemediationCreate,
    RiskAcceptanceCreate,
    RiskAcceptanceDecision,
    SeverityUpdate,
    SourceCreate,
    SourcePublic,
    SourceUpdate,
    TokenResponse,
    TransitionRequest,
    UserPublic,
    VerificationCreate,
)
from .security import AdminUser, CurrentUser, authenticate, create_access_token
from .seed import seed_database
from .services import SEVERITY_PRIORITY, SEVERITY_SCORE, allowed_actions, audit, finding_event, ingest_records, transition_finding


DEFAULT_GOVERNANCE_SETTINGS = {
    "severity_mapping": ("等级映射", {"critical": "严重", "high": "高危", "medium": "中危", "low": "低危", "info": "提示"}),
    "deduplication": ("去重规则", {"stable_id_enabled": True, "field_hash_enabled": True, "scope": "source_asset"}),
    "sla": ("SLA 策略", {"critical_days": 3, "high_days": 7, "medium_days": 30, "low_days": 60, "info_days": 90, "remind_before_days": 3}),
    "auto_assignment": ("自动分派", {"enabled": True, "strategy": "asset_owner", "fallback_to_admin": True}),
    "notifications": ("通知规则", {"assignment": True, "approaching_sla": True, "overdue": True, "verification_rejected": True, "channel": "in_app"}),
    "risk_acceptance": ("风险接受", {"max_days": 90, "require_compensating_control": True, "restore_on_expiry": True, "approver_role": "platform_admin"}),
}


@asynccontextmanager
async def lifespan(_: FastAPI):
    settings.storage_dir.mkdir(parents=True, exist_ok=True)
    create_schema()
    with SessionLocal() as session:
        seed_database(session)
    yield


app = FastAPI(title=settings.app_name, version="0.1.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.allowed_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def request_context(request: Request, call_next):
    request_id = request.headers.get("X-Request-ID") or str(uuid.uuid4())
    response = await call_next(request)
    response.headers["X-Request-ID"] = request_id
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    return response


def finding_query():
    return select(Finding).options(
        selectinload(Finding.asset).selectinload(Asset.owner),
        selectinload(Finding.owner),
        selectinload(Finding.assignee),
        selectinload(Finding.verifier),
    )


def finding_public(finding: Finding, user: User, detail: bool = False):
    schema = FindingDetail if detail else FindingSummary
    return schema.model_validate(finding).model_copy(update={"allowed_actions": allowed_actions(user, finding)})


def scoped_findings(stmt, user: User):
    if Role.PLATFORM_ADMIN.value in user.roles:
        return stmt
    scopes = []
    if Role.REMEDIATOR.value in user.roles:
        scopes.append(Finding.assignee_id == user.id)
    if Role.VERIFIER.value in user.roles:
        scopes.append(Finding.verifier_id == user.id)
    return stmt.where(or_(*scopes)) if scopes else stmt.where(False)


def get_accessible_finding(session: Session, user: User, finding_id: str) -> Finding:
    finding = session.scalar(scoped_findings(finding_query().where(Finding.id == finding_id), user))
    if not finding:
        raise HTTPException(status_code=404, detail={"code": "NOT_FOUND", "message": "风险不存在或无权访问"})
    return finding


def ensure_version(finding: Finding, version: int) -> None:
    if finding.version != version:
        raise HTTPException(status_code=409, detail={"code": "VERSION_CONFLICT", "message": "风险已被更新，请刷新后重试"})


def next_batch_no(session: Session) -> str:
    count = session.scalar(select(func.count()).select_from(ImportBatch)) or 0
    return f"IMP-{datetime.now().strftime('%y%m%d')}-{count + 1:04d}"


@app.get("/health")
def health():
    return {"status": "ok", "service": settings.app_name}


@app.post("/api/v1/auth/login", response_model=TokenResponse)
def login(payload: LoginRequest, session: Annotated[Session, Depends(get_session)]):
    user = authenticate(session, payload.username, payload.password)
    if not user:
        raise HTTPException(status_code=401, detail={"code": "INVALID_CREDENTIALS", "message": "用户名或密码错误"})
    audit(session, user, "LOGIN", "user", user.id)
    session.commit()
    return TokenResponse(access_token=create_access_token(user))


@app.get("/api/v1/me", response_model=UserPublic)
def me(user: CurrentUser):
    return user


@app.get("/api/v1/users", response_model=list[UserPublic])
def users(_: AdminUser, session: Annotated[Session, Depends(get_session)]):
    return session.scalars(select(User).where(User.enabled.is_(True)).order_by(User.display_name)).all()


@app.get("/api/v1/assets", response_model=list[AssetPublic])
def list_assets(user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    stmt = select(Asset).options(selectinload(Asset.owner)).order_by(Asset.name)
    if Role.PLATFORM_ADMIN.value not in user.roles:
        stmt = stmt.where(Asset.owner_id == user.id)
    return session.scalars(stmt).all()


@app.post("/api/v1/assets", response_model=AssetPublic, status_code=201)
def create_asset(payload: AssetCreate, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    asset = Asset(**payload.model_dump())
    session.add(asset)
    try:
        session.flush()
    except IntegrityError as exc:
        session.rollback()
        raise HTTPException(status_code=409, detail={"code": "ASSET_EXISTS", "message": "资产编码已存在"}) from exc
    audit(session, user, "ASSET_CREATED", "asset", asset.id, after={"asset_code": asset.asset_code, "name": asset.name})
    session.commit()
    return session.scalar(select(Asset).options(selectinload(Asset.owner)).where(Asset.id == asset.id))


@app.patch("/api/v1/assets/{asset_id}", response_model=AssetPublic)
def update_asset(asset_id: str, payload: AssetUpdate, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    asset = session.get(Asset, asset_id)
    if not asset:
        raise HTTPException(status_code=404, detail={"code": "ASSET_NOT_FOUND", "message": "资产不存在"})
    if payload.owner_id and not session.get(User, payload.owner_id):
        raise HTTPException(status_code=422, detail={"code": "OWNER_NOT_FOUND", "message": "资产 Owner 不存在"})
    before = {field: getattr(asset, field) for field in type(payload).model_fields}
    for field, value in payload.model_dump().items():
        setattr(asset, field, value)
    audit(session, user, "ASSET_UPDATED", "asset", asset.id, before, payload.model_dump(mode="json"))
    session.commit()
    return session.scalar(select(Asset).options(selectinload(Asset.owner)).where(Asset.id == asset.id))


@app.get("/api/v1/sources", response_model=list[SourcePublic])
def list_sources(_: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    return session.scalars(select(Source).order_by(Source.name)).all()


@app.post("/api/v1/sources", response_model=SourcePublic, status_code=201)
def create_source(payload: SourceCreate, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    source = Source(**payload.model_dump())
    session.add(source)
    try:
        session.flush()
    except IntegrityError as exc:
        session.rollback()
        raise HTTPException(status_code=409, detail={"code": "SOURCE_EXISTS", "message": "来源编码已存在"}) from exc
    audit(session, user, "SOURCE_CREATED", "source", source.id, after={"source_code": source.source_code, "name": source.name})
    session.commit()
    return source


@app.patch("/api/v1/sources/{source_id}", response_model=SourcePublic)
def update_source(source_id: str, payload: SourceUpdate, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    source = session.get(Source, source_id)
    if not source:
        raise HTTPException(status_code=404, detail={"code": "SOURCE_NOT_FOUND", "message": "来源不存在"})
    before = {field: getattr(source, field) for field in type(payload).model_fields}
    for field, value in payload.model_dump().items():
        setattr(source, field, value)
    audit(session, user, "SOURCE_UPDATED", "source", source.id, before, payload.model_dump(mode="json"))
    session.commit()
    session.refresh(source)
    return source


def ensure_governance_settings(session: Session) -> list[GovernanceSetting]:
    existing = {item.key: item for item in session.scalars(select(GovernanceSetting)).all()}
    for key, (title, config) in DEFAULT_GOVERNANCE_SETTINGS.items():
        if key not in existing:
            setting = GovernanceSetting(key=key, title=title, config=config)
            session.add(setting)
            existing[key] = setting
    session.commit()
    return [existing[key] for key in DEFAULT_GOVERNANCE_SETTINGS]


@app.get("/api/v1/governance-settings", response_model=list[GovernanceSettingPublic])
def list_governance_settings(_: AdminUser, session: Annotated[Session, Depends(get_session)]):
    return ensure_governance_settings(session)


@app.patch("/api/v1/governance-settings/{setting_key}", response_model=GovernanceSettingPublic)
def update_governance_setting(setting_key: str, payload: GovernanceSettingUpdate, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    ensure_governance_settings(session)
    setting = session.scalar(select(GovernanceSetting).where(GovernanceSetting.key == setting_key))
    if not setting:
        raise HTTPException(status_code=404, detail={"code": "SETTING_NOT_FOUND", "message": "治理配置不存在"})
    if setting.version != payload.version:
        raise HTTPException(status_code=409, detail={"code": "VERSION_CONFLICT", "message": "配置已被更新，请刷新后重试"})
    before = dict(setting.config)
    setting.config = payload.config
    setting.version += 1
    audit(session, user, "GOVERNANCE_SETTING_UPDATED", "governance_setting", setting.key, before, setting.config)
    session.commit()
    session.refresh(setting)
    return setting


@app.get("/api/v1/import-batches", response_model=list[BatchPublic])
def list_batches(_: AdminUser, session: Annotated[Session, Depends(get_session)]):
    return session.scalars(select(ImportBatch).order_by(ImportBatch.created_at.desc()).limit(100)).all()


@app.get("/api/v1/import-batches/template")
def download_import_template(_: AdminUser):
    headers = ["source_finding_id", "source_rule_id", "title", "description", "recommendation", "severity", "asset_code", "asset_external_id", "asset_name", "asset_type", "team", "location", "observed_at"]
    descriptions = {
        "source_finding_id": "来源系统中的风险唯一 ID，推荐填写，用于稳定去重",
        "source_rule_id": "来源规则或检查项 ID",
        "title": "风险标题，必填",
        "description": "风险描述",
        "recommendation": "整改建议",
        "severity": "风险等级，必填：critical / high / medium / low / info",
        "asset_code": "平台资产编码；与 asset_external_id 至少填写一个",
        "asset_external_id": "来源系统中的资产 ID",
        "asset_name": "资产名称，新资产自动创建时使用",
        "asset_type": "资产类型，如 application、api、database、cloud",
        "team": "所属团队",
        "location": "发现位置，如 URL、文件路径、云资源路径",
        "observed_at": "发现时间，ISO 8601 格式，例如 2026-09-05T10:30:00+08:00",
    }
    example = ["SCAN-2026-001", "WEAK_PASSWORD", "公网管理后台存在弱口令策略", "公网入口允许弱口令，可能导致未授权访问。", "关闭公网入口并启用 MFA。", "critical", "OPS-PLATFORM", "", "统一运维平台", "application", "基础架构组", "https://ops.example.com/admin", "2026-09-05T10:30:00+08:00"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "风险导入"
    sheet.append(headers)
    sheet.append(example)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M2"
    header_fill = PatternFill("solid", fgColor="2478D4")
    required_fill = PatternFill("solid", fgColor="FFF2CC")
    for index, header in enumerate(headers, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.comment = Comment(descriptions[header], "RiskHub")
        sheet.column_dimensions[cell.column_letter].width = max(18, min(42, len(descriptions[header]) * 1.4))
        if header in {"title", "severity"}:
            sheet.cell(2, index).fill = required_fill
    severity_validation = DataValidation(type="list", formula1='"critical,high,medium,low,info"', allow_blank=False)
    severity_validation.error = "请选择 critical、high、medium、low 或 info"
    severity_validation.errorTitle = "无效风险等级"
    sheet.add_data_validation(severity_validation)
    severity_validation.add("F2:F10001")
    required_rule = FormulaRule(formula=['OR($C2="",$F2="")'], fill=PatternFill("solid", fgColor="FCE8E6"))
    sheet.conditional_formatting.add("A2:M10001", required_rule)
    guide = workbook.create_sheet("填写说明")
    guide.append(["字段", "是否必填", "填写说明"])
    for header in headers:
        guide.append([header, "是" if header in {"title", "severity"} else "否", descriptions[header]])
    guide.freeze_panes = "A2"
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 72
    for cell in guide[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="riskhub-finding-import-template.xlsx"'},
    )


@app.post("/api/v1/import-batches/api", response_model=BatchPublic, status_code=201)
def api_import(
    payload: APIImportRequest,
    user: AdminUser,
    session: Annotated[Session, Depends(get_session)],
    idempotency_key: Annotated[str, Header(alias="Idempotency-Key")],
):
    existing = session.scalar(select(ImportBatch).where(ImportBatch.source_id == payload.source_id, ImportBatch.idempotency_key == idempotency_key))
    if existing:
        return existing
    source = session.get(Source, payload.source_id)
    if not source or not source.enabled:
        raise HTTPException(status_code=404, detail={"code": "SOURCE_NOT_FOUND", "message": "来源不存在或已停用"})
    batch = ImportBatch(batch_no=next_batch_no(session), source_id=source.id, idempotency_key=idempotency_key, status=BatchStatus.PENDING, filename="API 请求")
    session.add(batch)
    session.flush()
    return ingest_records(session, batch, source, payload.records, user)


def excel_records(content: bytes) -> list[ObservationInput]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows, [])]
    required = {"title", "severity"}
    if not required.issubset(headers):
        raise HTTPException(status_code=422, detail={"code": "INVALID_TEMPLATE", "message": "Excel 必须包含 title 和 severity 列"})
    records: list[ObservationInput] = []
    for row in rows:
        data = {headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
        if not any(value is not None and str(value).strip() for value in data.values()):
            continue
        records.append(ObservationInput(**data))
        if len(records) > 10_000:
            raise HTTPException(status_code=413, detail={"code": "TOO_MANY_RECORDS", "message": "单批次最多 10,000 条记录"})
    return records


@app.post("/api/v1/import-batches/files", response_model=BatchPublic, status_code=201)
async def file_import(
    user: AdminUser,
    session: Annotated[Session, Depends(get_session)],
    source_id: str,
    file: Annotated[UploadFile, File()],
    idempotency_key: Annotated[str | None, Header(alias="Idempotency-Key")] = None,
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=415, detail={"code": "INVALID_FILE_TYPE", "message": "仅支持 .xlsx 文件"})
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail={"code": "FILE_TOO_LARGE", "message": "文件不能超过 10 MB"})
    source = session.get(Source, source_id)
    if not source or not source.enabled:
        raise HTTPException(status_code=404, detail={"code": "SOURCE_NOT_FOUND", "message": "来源不存在或已停用"})
    digest = hashlib.sha256(content).hexdigest()
    key = idempotency_key or digest
    existing = session.scalar(select(ImportBatch).where(ImportBatch.source_id == source_id, ImportBatch.idempotency_key == key))
    if existing:
        return existing
    safe_name = f"{uuid.uuid4()}.xlsx"
    (settings.storage_dir / safe_name).write_bytes(content)
    records = excel_records(content)
    batch = ImportBatch(batch_no=next_batch_no(session), source_id=source.id, idempotency_key=key, filename=Path(file.filename).name, file_sha256=digest, status=BatchStatus.PENDING)
    session.add(batch)
    session.flush()
    return ingest_records(session, batch, source, records, user)


@app.get("/api/v1/findings")
def list_findings(
    user: CurrentUser,
    session: Annotated[Session, Depends(get_session)],
    q: str | None = None,
    severity: Severity | None = None,
    status: FindingStatus | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
):
    stmt = scoped_findings(finding_query(), user)
    count_stmt = scoped_findings(select(func.count()).select_from(Finding), user)
    filters = []
    if q:
        filters.append(or_(Finding.finding_no.ilike(f"%{q}%"), Finding.title.ilike(f"%{q}%")))
    if severity:
        filters.append(Finding.severity == severity)
    if status:
        filters.append(Finding.status == status)
    if filters:
        stmt = stmt.where(*filters)
        count_stmt = count_stmt.where(*filters)
    total = session.scalar(count_stmt) or 0
    items = session.scalars(stmt.order_by(Finding.last_seen_at.desc()).offset((page - 1) * page_size).limit(page_size)).all()
    return {"items": [finding_public(item, user) for item in items], "total": total, "page": page, "page_size": page_size}


@app.get("/api/v1/findings/{finding_id}", response_model=FindingDetail)
def get_finding(finding_id: str, user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    return finding_public(get_accessible_finding(session, user, finding_id), user, detail=True)


@app.patch("/api/v1/findings/{finding_id}", response_model=FindingDetail)
def update_finding(finding_id: str, payload: FindingUpdate, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    ensure_version(finding, payload.version)
    asset = session.get(Asset, payload.asset_id)
    if not asset:
        raise HTTPException(status_code=422, detail={"code": "ASSET_NOT_FOUND", "message": "关联资产不存在"})
    before = {
        "title": finding.title, "category": finding.category, "description": finding.description,
        "recommendation": finding.recommendation, "severity": finding.severity.value, "asset_id": finding.asset_id,
    }
    finding.title = payload.title
    finding.category = payload.category
    finding.description = payload.description
    finding.recommendation = payload.recommendation
    finding.severity = payload.severity
    finding.severity_overridden = True
    finding.risk_score = SEVERITY_SCORE[payload.severity]
    finding.priority = SEVERITY_PRIORITY[payload.severity]
    if finding.asset_id != payload.asset_id:
        finding.asset = asset
        session.execute(update(Observation).where(Observation.finding_id == finding.id).values(asset_id=payload.asset_id))
    finding.version += 1
    after = {**payload.model_dump(mode="json"), "version": finding.version}
    audit(session, user, "FINDING_UPDATED", "finding", finding.id, before, after)
    finding_event(session, finding, "FINDING_UPDATED", user, payload={"reason": payload.reason})
    session.commit()
    return finding_public(get_accessible_finding(session, user, finding_id), user, detail=True)


@app.delete("/api/v1/findings/{finding_id}", status_code=204)
def delete_finding(finding_id: str, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    before = {"finding_no": finding.finding_no, "title": finding.title, "status": finding.status.value}
    session.execute(delete(Verification).where(Verification.finding_id == finding.id))
    session.execute(delete(Remediation).where(Remediation.finding_id == finding.id))
    session.execute(delete(RiskAcceptance).where(RiskAcceptance.finding_id == finding.id))
    session.execute(delete(Observation).where(Observation.finding_id == finding.id))
    session.execute(delete(FindingEvent).where(FindingEvent.finding_id == finding.id))
    session.delete(finding)
    audit(session, user, "FINDING_DELETED", "finding", finding.id, before, None)
    session.commit()


@app.get("/api/v1/findings/{finding_id}/observations", response_model=list[ObservationPublic])
def finding_observations(finding_id: str, user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    get_accessible_finding(session, user, finding_id)
    return session.scalars(select(Observation).options(selectinload(Observation.source)).where(Observation.finding_id == finding_id).order_by(Observation.observed_at.desc())).all()


@app.get("/api/v1/findings/{finding_id}/events", response_model=list[EventPublic])
def finding_events(finding_id: str, user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    get_accessible_finding(session, user, finding_id)
    return session.scalars(select(FindingEvent).options(selectinload(FindingEvent.actor)).where(FindingEvent.finding_id == finding_id).order_by(FindingEvent.occurred_at.desc())).all()


@app.patch("/api/v1/findings/{finding_id}/assignment", response_model=FindingDetail)
def update_assignment(finding_id: str, payload: AssignmentUpdate, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    ensure_version(finding, payload.version)
    if payload.assignee_id == payload.verifier_id:
        raise HTTPException(status_code=422, detail={"code": "DUTY_CONFLICT", "message": "整改人与验证人不能相同"})
    users = session.scalars(select(User).where(User.id.in_([payload.owner_id, payload.assignee_id, payload.verifier_id]))).all()
    if len(users) != 3:
        raise HTTPException(status_code=422, detail={"code": "INVALID_USERS", "message": "分派用户不存在"})
    before = {"owner_id": finding.owner_id, "assignee_id": finding.assignee_id, "verifier_id": finding.verifier_id, "due_at": finding.due_at.isoformat() if finding.due_at else None}
    finding.owner_id, finding.assignee_id, finding.verifier_id, finding.due_at = payload.owner_id, payload.assignee_id, payload.verifier_id, payload.due_at
    finding.version += 1
    audit(session, user, "ASSIGNMENT_UPDATED", "finding", finding.id, before, {**payload.model_dump(mode="json"), "version": finding.version})
    finding_event(session, finding, "ASSIGNMENT_UPDATED", user, payload={"reason": payload.reason})
    session.commit()
    finding = get_accessible_finding(session, user, finding_id)
    return finding_public(finding, user, detail=True)


@app.patch("/api/v1/findings/{finding_id}/severity", response_model=FindingDetail)
def update_severity(finding_id: str, payload: SeverityUpdate, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    ensure_version(finding, payload.version)
    before = finding.severity
    finding.severity = payload.severity
    finding.severity_overridden = True
    finding.version += 1
    finding_event(session, finding, "SEVERITY_UPDATED", user, payload={"from": before.value, "to": payload.severity.value, "reason": payload.reason})
    audit(session, user, "SEVERITY_UPDATED", "finding", finding.id, {"severity": before.value}, {"severity": payload.severity.value, "reason": payload.reason})
    session.commit()
    return finding_public(get_accessible_finding(session, user, finding_id), user, detail=True)


@app.post("/api/v1/findings/{finding_id}/transitions", response_model=FindingDetail)
def transition(finding_id: str, payload: TransitionRequest, user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    transition_finding(session, finding, user, payload.action, payload.reason, payload.version)
    return finding_public(get_accessible_finding(session, user, finding_id), user, detail=True)


@app.post("/api/v1/findings/{finding_id}/remediations", response_model=FindingDetail)
def submit_remediation(finding_id: str, payload: RemediationCreate, user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    ensure_version(finding, payload.version)
    is_admin = Role.PLATFORM_ADMIN.value in user.roles
    if finding.status != FindingStatus.IN_REMEDIATION:
        raise HTTPException(status_code=409, detail={"code": "INVALID_STATUS", "message": "仅整改中的风险可以提交整改结果"})
    if not is_admin and (Role.REMEDIATOR.value not in user.roles or finding.assignee_id != user.id):
        raise HTTPException(status_code=403, detail={"code": "FORBIDDEN", "message": "仅当前整改人可以提交"})
    revision = (session.scalar(select(func.count()).select_from(Remediation).where(Remediation.finding_id == finding.id)) or 0) + 1
    remediation = Remediation(finding_id=finding.id, submitted_by=user.id, description=payload.description, evidence=payload.evidence, revision=revision)
    session.add(remediation)
    before = finding.status
    finding.status = FindingStatus.PENDING_VERIFICATION
    finding.resolved_at = utcnow()
    finding.version += 1
    finding_event(session, finding, "REMEDIATION_SUBMITTED", user, before, finding.status, {"revision": revision})
    audit(session, user, "REMEDIATION_SUBMITTED", "finding", finding.id, {"status": before.value}, {"status": finding.status.value, "revision": revision})
    session.commit()
    return finding_public(get_accessible_finding(session, user, finding_id), user, detail=True)


@app.post("/api/v1/findings/{finding_id}/verifications", response_model=FindingDetail)
def verify_finding(finding_id: str, payload: VerificationCreate, user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    ensure_version(finding, payload.version)
    is_admin = Role.PLATFORM_ADMIN.value in user.roles
    if finding.status != FindingStatus.PENDING_VERIFICATION:
        raise HTTPException(status_code=409, detail={"code": "INVALID_STATUS", "message": "仅待验证风险可以执行验证"})
    if not is_admin and (Role.VERIFIER.value not in user.roles or finding.verifier_id != user.id):
        raise HTTPException(status_code=403, detail={"code": "FORBIDDEN", "message": "仅当前验证人可以验证"})
    if not is_admin and finding.assignee_id == user.id:
        raise HTTPException(status_code=422, detail={"code": "DUTY_CONFLICT", "message": "整改人与验证人不能为同一人"})
    remediation = session.scalar(select(Remediation).where(Remediation.finding_id == finding.id).order_by(Remediation.revision.desc()))
    if not remediation:
        raise HTTPException(status_code=422, detail={"code": "REMEDIATION_REQUIRED", "message": "没有可验证的整改记录"})
    session.add(Verification(finding_id=finding.id, remediation_id=remediation.id, verified_by=user.id, result=payload.result, method=payload.method, comment=payload.comment))
    before = finding.status
    if payload.result == "passed":
        finding.status = FindingStatus.CLOSED
        finding.verified_at = utcnow()
        finding.closed_at = finding.verified_at
    else:
        finding.status = FindingStatus.IN_REMEDIATION
        finding.resolved_at = None
    finding.version += 1
    event_type = "VERIFICATION_PASSED" if payload.result == "passed" else "VERIFICATION_REJECTED"
    finding_event(session, finding, event_type, user, before, finding.status, {"method": payload.method, "comment": payload.comment})
    audit(session, user, event_type, "finding", finding.id, {"status": before.value}, {"status": finding.status.value})
    session.commit()
    return finding_public(get_accessible_finding(session, user, finding_id), user, detail=True)


@app.post("/api/v1/findings/{finding_id}/risk-acceptances", response_model=FindingDetail)
def request_acceptance(finding_id: str, payload: RiskAcceptanceCreate, user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    ensure_version(finding, payload.version)
    is_admin = Role.PLATFORM_ADMIN.value in user.roles
    if finding.status not in {FindingStatus.PENDING_REMEDIATION, FindingStatus.IN_REMEDIATION}:
        raise HTTPException(status_code=409, detail={"code": "INVALID_STATUS", "message": "当前状态不能申请风险接受"})
    if not is_admin and (Role.REMEDIATOR.value not in user.roles or finding.assignee_id != user.id):
        raise HTTPException(status_code=403, detail={"code": "FORBIDDEN", "message": "仅当前整改人可以申请"})
    if payload.expires_at <= datetime.now(timezone.utc):
        raise HTTPException(status_code=422, detail={"code": "INVALID_EXPIRY", "message": "到期时间必须晚于当前时间"})
    session.add(RiskAcceptance(finding_id=finding.id, reason=payload.reason, compensating_control=payload.compensating_control, requested_by=user.id, expires_at=payload.expires_at))
    before = finding.status
    finding.status = FindingStatus.ACCEPTANCE_REQUESTED
    finding.version += 1
    finding_event(session, finding, "RISK_ACCEPTANCE_REQUESTED", user, before, finding.status, {"expires_at": payload.expires_at.isoformat()})
    audit(session, user, "RISK_ACCEPTANCE_REQUESTED", "finding", finding.id, {"status": before.value}, {"status": finding.status.value})
    session.commit()
    return finding_public(get_accessible_finding(session, user, finding_id), user, detail=True)


@app.post("/api/v1/findings/{finding_id}/risk-acceptances/decision", response_model=FindingDetail)
def decide_acceptance(finding_id: str, payload: RiskAcceptanceDecision, user: AdminUser, session: Annotated[Session, Depends(get_session)]):
    finding = get_accessible_finding(session, user, finding_id)
    ensure_version(finding, payload.version)
    if finding.status != FindingStatus.ACCEPTANCE_REQUESTED:
        raise HTTPException(status_code=409, detail={"code": "INVALID_STATUS", "message": "当前没有待审批的风险接受"})
    acceptance = session.scalar(select(RiskAcceptance).where(RiskAcceptance.finding_id == finding.id, RiskAcceptance.status == "requested").order_by(RiskAcceptance.expires_at.desc()))
    if not acceptance:
        raise HTTPException(status_code=404, detail={"code": "ACCEPTANCE_NOT_FOUND", "message": "风险接受申请不存在"})
    if acceptance.requested_by == user.id:
        raise HTTPException(status_code=422, detail={"code": "SELF_APPROVAL", "message": "申请人不能审批自己的申请"})
    acceptance.approved_by = user.id
    acceptance.decision_comment = payload.comment
    acceptance.status = "approved" if payload.approved else "rejected"
    acceptance.starts_at = utcnow() if payload.approved else None
    before = finding.status
    finding.status = FindingStatus.RISK_ACCEPTED if payload.approved else FindingStatus.PENDING_REMEDIATION
    finding.version += 1
    event_type = "RISK_ACCEPTANCE_APPROVED" if payload.approved else "RISK_ACCEPTANCE_REJECTED"
    finding_event(session, finding, event_type, user, before, finding.status, {"comment": payload.comment})
    audit(session, user, event_type, "finding", finding.id, {"status": before.value}, {"status": finding.status.value})
    session.commit()
    return finding_public(get_accessible_finding(session, user, finding_id), user, detail=True)


@app.get("/api/v1/dashboard/summary")
def dashboard_summary(user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    findings = session.scalars(scoped_findings(select(Finding), user)).all()
    now = datetime.now(timezone.utc)
    active = [item for item in findings if item.status not in {FindingStatus.CLOSED, FindingStatus.FALSE_POSITIVE}]
    approaching_cutoff = now + timedelta(days=3)

    def as_utc(value: datetime) -> datetime:
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)

    return {
        "active": len(active),
        "critical_high": sum(item.severity in {Severity.CRITICAL, Severity.HIGH} for item in active),
        "sla_approaching": sum(bool(item.due_at and now <= as_utc(item.due_at) <= approaching_cutoff) for item in active),
        "closed": sum(item.status == FindingStatus.CLOSED for item in findings),
        "by_status": {status.value: sum(item.status == status for item in findings) for status in FindingStatus},
        "by_severity": {severity.value: sum(item.severity == severity for item in active) for severity in Severity},
    }


@app.get("/api/v1/dashboard/priorities", response_model=list[FindingSummary])
def dashboard_priorities(user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    stmt = scoped_findings(finding_query().where(Finding.status.not_in([FindingStatus.CLOSED, FindingStatus.FALSE_POSITIVE])), user)
    items = session.scalars(stmt.order_by(Finding.risk_score.desc(), Finding.due_at.asc()).limit(8)).all()
    return [finding_public(item, user) for item in items]


@app.get("/api/v1/reports/overview")
def report_overview(user: CurrentUser, session: Annotated[Session, Depends(get_session)]):
    findings = session.scalars(scoped_findings(select(Finding), user)).all()
    batches = session.scalars(select(ImportBatch)).all() if Role.PLATFORM_ADMIN.value in user.roles else []
    closed = [item for item in findings if item.closed_at]
    on_time = [item for item in closed if not item.due_at or item.closed_at <= item.due_at]
    return {
        "total_findings": len(findings),
        "active_findings": sum(item.status not in {FindingStatus.CLOSED, FindingStatus.FALSE_POSITIVE} for item in findings),
        "sla_compliance": round(len(on_time) / len(closed) * 100, 1) if closed else 100.0,
        "import_success_rate": round(sum(item.success_count for item in batches) / max(sum(item.total_count for item in batches), 1) * 100, 1),
        "reopened": sum(item.reopen_count for item in findings),
    }


@app.get("/api/v1/audit-events", response_model=list[AuditPublic])
def list_audit_events(user: AdminUser, session: Annotated[Session, Depends(get_session)], limit: int = Query(100, ge=1, le=500)):
    return session.scalars(select(AuditEvent).options(selectinload(AuditEvent.actor)).order_by(AuditEvent.occurred_at.desc()).limit(limit)).all()
