import io
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook, load_workbook


def first_with_status(client, headers, status):
    response = client.get(f"/api/v1/findings?status={status}", headers=headers)
    assert response.status_code == 200, response.text
    items = response.json()["items"]
    assert items, f"no finding with status {status}"
    return items[0]


def test_health_login_and_current_user(client, admin_headers):
    assert client.get("/health").json()["status"] == "ok"
    response = client.get("/api/v1/me", headers=admin_headers)
    assert response.status_code == 200
    assert response.json()["roles"] == ["platform_admin"]
    assert response.headers["x-request-id"]
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"


def test_role_scoped_access_and_admin_only_endpoint(client, admin_headers, remediator_headers):
    admin_findings = client.get("/api/v1/findings", headers=admin_headers).json()
    remediation_findings = client.get("/api/v1/findings", headers=remediator_headers).json()
    assert admin_findings["total"] >= 6
    assert 0 < remediation_findings["total"] <= admin_findings["total"]
    assert client.get("/api/v1/audit-events", headers=remediator_headers).status_code == 403
    target = remediation_findings["items"][0]
    response = client.patch(
        f"/api/v1/findings/{target['id']}/severity",
        headers=remediator_headers,
        json={"severity": "low", "reason": "越权测试", "version": target["version"]},
    )
    assert response.status_code == 403


def test_admin_syncs_feishu_directory_and_dispatches_risk(client, admin_headers, remediator_headers, monkeypatch):
    from riskhub import main

    monkeypatch.setattr(main.settings, "feishu_app_id", "test-app")
    monkeypatch.setattr(main.settings, "feishu_app_secret", "test-secret")
    monkeypatch.setattr(main.settings, "feishu_department_id", "od-test-sre")
    monkeypatch.setattr(main.settings, "feishu_department_ids", "")
    monkeypatch.setattr(main.settings, "feishu_department_name", "SRE")
    monkeypatch.setattr(main.settings, "feishu_https_proxy", "http://proxy.internal:8080")
    monkeypatch.setattr(main, "list_department_users", lambda: [
        {"open_id": "ou-sre-1", "name": "飞书整改人员", "status": {"is_activated": True}},
        {"open_id": "ou-sre-2", "name": "飞书验证人员", "status": {"is_activated": True}},
    ])
    sent_messages = []
    monkeypatch.setattr(main, "send_assignment_message", lambda open_id, **payload: sent_messages.append((open_id, payload)))

    assert client.post("/api/v1/integrations/feishu/sync", headers=remediator_headers).status_code == 403
    synced = client.post("/api/v1/integrations/feishu/sync", headers=admin_headers)
    assert synced.status_code == 200, synced.text
    assert synced.json() == {"created_count": 2, "updated_count": 0, "disabled_count": 0, "total_count": 2}
    status = client.get("/api/v1/integrations/feishu", headers=admin_headers).json()
    assert status["configured"] is True
    assert status["proxy_configured"] is True
    assert status["app_id_hint"] == "test…st-app"
    assert status["department_name"] == "SRE"
    assert status["department_count"] == 1
    assert status["member_count"] == 2

    users = client.get("/api/v1/users", headers=admin_headers).json()
    assignee = next(item for item in users if item["display_name"] == "飞书整改人员")
    verifier = next(item for item in users if item["username"] == "verifier")
    admin = next(item for item in users if item["username"] == "admin")
    finding = first_with_status(client, admin_headers, "pending_confirmation")
    assigned = client.patch(
        f"/api/v1/findings/{finding['id']}/assignment",
        headers=admin_headers,
        json={
            "owner_id": admin["id"], "assignee_id": assignee["id"], "verifier_id": verifier["id"],
            "due_at": (datetime.now(timezone.utc) + timedelta(days=3)).isoformat(),
            "reason": "派发给 SRE", "version": finding["version"],
        },
    )
    assert assigned.status_code == 200, assigned.text
    assert sent_messages[0][0] == "ou-sre-1"
    assert sent_messages[0][1]["finding_id"] == finding["id"]


def test_admin_can_edit_and_delete_finding(client, admin_headers, remediator_headers):
    source = client.get("/api/v1/sources", headers=admin_headers).json()[0]
    created_batch = client.post(
        "/api/v1/import-batches/api",
        headers={**admin_headers, "Idempotency-Key": "edit-delete-finding-test"},
        json={"source_id": source["id"], "records": [{
            "source_finding_id": "EDIT-DELETE-001", "source_rule_id": "EDIT-DELETE",
            "title": "待修正的风险标题", "description": "原始描述", "recommendation": "原始建议",
            "severity": "medium", "asset_code": "OPS-PLATFORM", "location": "/edit-delete-test",
        }]},
    )
    assert created_batch.status_code == 201, created_batch.text
    finding = client.get("/api/v1/findings?q=待修正的风险标题", headers=admin_headers).json()["items"][0]
    assets = client.get("/api/v1/assets", headers=admin_headers).json()
    new_asset = next(item for item in assets if item["id"] != finding["asset"]["id"])
    payload = {
        "title": "修正后的风险标题", "category": "configuration", "description": "修正后的描述",
        "recommendation": "修正后的整改建议", "severity": "high", "asset_id": new_asset["id"],
        "version": finding["version"], "reason": "修正错误导入信息",
    }
    assert client.patch(f"/api/v1/findings/{finding['id']}", headers=remediator_headers, json=payload).status_code == 403
    updated = client.patch(f"/api/v1/findings/{finding['id']}", headers=admin_headers, json=payload)
    assert updated.status_code == 200, updated.text
    assert updated.json()["title"] == "修正后的风险标题"
    assert updated.json()["severity"] == "high"
    assert updated.json()["risk_score"] == 8.0
    assert updated.json()["priority"] == "P1"
    assert updated.json()["asset"]["id"] == new_asset["id"]

    assert client.delete(f"/api/v1/findings/{finding['id']}", headers=remediator_headers).status_code == 403
    deleted = client.delete(f"/api/v1/findings/{finding['id']}", headers=admin_headers)
    assert deleted.status_code == 204
    assert client.get(f"/api/v1/findings/{finding['id']}", headers=admin_headers).status_code == 404
    actions = {item["action"] for item in client.get("/api/v1/audit-events", headers=admin_headers).json()}
    assert {"FINDING_UPDATED", "FINDING_DELETED"}.issubset(actions)


def test_findings_pagination_and_admin_bulk_delete(client, admin_headers, remediator_headers):
    first_page = client.get("/api/v1/findings?page=1&page_size=2", headers=admin_headers)
    second_page = client.get("/api/v1/findings?page=2&page_size=2", headers=admin_headers)
    assert first_page.status_code == 200
    assert second_page.status_code == 200
    assert first_page.json()["page"] == 1
    assert second_page.json()["page"] == 2
    assert first_page.json()["page_size"] == 2
    first_ids = [item["id"] for item in first_page.json()["items"]]
    second_ids = [item["id"] for item in second_page.json()["items"]]
    assert len(first_ids) == 2
    assert set(first_ids).isdisjoint(second_ids)

    assert client.post("/api/v1/findings/bulk-delete", headers=remediator_headers, json={"ids": first_ids}).status_code == 403
    deleted = client.post("/api/v1/findings/bulk-delete", headers=admin_headers, json={"ids": first_ids})
    assert deleted.status_code == 200, deleted.text
    assert deleted.json() == {"deleted_count": 2}
    remaining_ids = {item["id"] for item in client.get("/api/v1/findings?page_size=100", headers=admin_headers).json()["items"]}
    assert not remaining_ids.intersection(first_ids)


def test_admin_can_edit_assets_and_source_configuration(client, admin_headers, remediator_headers):
    asset = client.get("/api/v1/assets", headers=admin_headers).json()[0]
    updated_asset = client.patch(
        f"/api/v1/assets/{asset['id']}",
        headers=admin_headers,
        json={
            "name": f"{asset['name']}（已编辑）",
            "type": asset["type"],
            "external_id": asset["external_id"],
            "business_system": asset["business_system"],
            "team": "平台治理组",
            "owner_id": asset["owner"]["id"],
            "importance": asset["importance"],
            "exposure": asset["exposure"],
            "environment": "production",
            "status": "active",
        },
    )
    assert updated_asset.status_code == 200, updated_asset.text
    assert updated_asset.json()["team"] == "平台治理组"
    assert client.patch(f"/api/v1/assets/{asset['id']}", headers=remediator_headers, json={}).status_code == 403

    source = client.get("/api/v1/sources", headers=admin_headers).json()[0]
    updated_source = client.patch(
        f"/api/v1/sources/{source['id']}",
        headers=admin_headers,
        json={
            "name": source["name"],
            "ingestion_type": source["ingestion_type"],
            "adapter_type": "custom-mapping",
            "enabled": True,
            "mapping_config": {"title": "issue_name", "severity": "risk_level", "asset": "asset_code", "location": "target"},
        },
    )
    assert updated_source.status_code == 200, updated_source.text
    assert updated_source.json()["mapping_config"]["severity"] == "risk_level"
    audit_events = client.get("/api/v1/audit-events", headers=admin_headers).json()
    assert {"ASSET_UPDATED", "SOURCE_UPDATED"}.issubset({item["action"] for item in audit_events})


def test_api_import_idempotency_and_deduplication(client, admin_headers):
    source = client.get("/api/v1/sources", headers=admin_headers).json()[0]
    payload = {
        "source_id": source["id"],
        "records": [{
            "source_finding_id": "IDEMPOTENT-001",
            "source_rule_id": "TEST-RULE",
            "title": "幂等与去重测试风险",
            "severity": "high",
            "asset_code": "OPS-PLATFORM",
            "location": "/admin"
        }]
    }
    headers = {**admin_headers, "Idempotency-Key": "api-test-idempotent"}
    first = client.post("/api/v1/import-batches/api", headers=headers, json=payload)
    second = client.post("/api/v1/import-batches/api", headers=headers, json=payload)
    assert first.status_code == 201, first.text
    assert second.status_code == 201, second.text
    assert first.json()["id"] == second.json()["id"]
    findings = client.get("/api/v1/findings?q=幂等与去重测试风险", headers=admin_headers).json()
    assert findings["total"] == 1
    assert findings["items"][0]["observation_count"] == 1


def test_custom_manual_risk_entry_creates_import_batch(client, admin_headers):
    source = client.get("/api/v1/sources", headers=admin_headers).json()[0]
    response = client.post(
        "/api/v1/import-batches/api",
        headers={**admin_headers, "Idempotency-Key": "manual-entry-test"},
        json={
            "source_id": source["id"],
            "records": [{
                "source_finding_id": "MANUAL-001",
                "title": "手工录入的风险",
                "description": "通过自定义录入表单提交",
                "recommendation": "完成验证后关闭",
                "severity": "medium",
                "asset_code": "USER-API",
                "location": "/api/manual",
            }],
        },
    )
    assert response.status_code == 201, response.text
    assert response.json()["filename"] == "API 请求"
    assert response.json()["success_count"] == 1
    finding = client.get("/api/v1/findings?q=手工录入的风险", headers=admin_headers).json()
    assert finding["total"] == 1


def test_excel_import(client, admin_headers):
    source = next(item for item in client.get("/api/v1/sources", headers=admin_headers).json() if item["ingestion_type"] == "excel")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["source_finding_id", "source_rule_id", "title", "severity", "asset_code", "location"])
    sheet.append(["EXCEL-001", "EXCEL-RULE", "Excel 导入测试风险", "medium", "MEMBER-DB", "member-db:5432"])
    content = io.BytesIO()
    workbook.save(content)
    response = client.post(
        f"/api/v1/import-batches/files?source_id={source['id']}",
        headers={**admin_headers, "Idempotency-Key": "excel-test"},
        files={"file": ("findings.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201, response.text
    assert response.json()["success_count"] == 1


def test_excel_import_template_has_guidance_and_validation(client, admin_headers):
    response = client.get("/api/v1/import-batches/template", headers=admin_headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["风险导入", "填写说明"]
    sheet = workbook["风险导入"]
    assert [cell.value for cell in sheet[1]][:6] == ["source_finding_id", "source_rule_id", "title", "description", "recommendation", "severity"]
    assert sheet["C1"].comment is not None
    assert len(sheet.data_validations.dataValidation) == 1
    assert "F2:F10001" in str(sheet.data_validations.dataValidation[0].sqref)


def test_governance_settings_can_be_updated_and_are_audited(client, admin_headers, remediator_headers):
    response = client.get("/api/v1/governance-settings", headers=admin_headers)
    assert response.status_code == 200
    assert len(response.json()) == 6
    sla = next(item for item in response.json() if item["key"] == "sla")
    config = {**sla["config"], "high_days": 10}
    updated = client.patch(
        "/api/v1/governance-settings/sla",
        headers=admin_headers,
        json={"config": config, "version": sla["version"]},
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["config"]["high_days"] == 10
    assert updated.json()["version"] == sla["version"] + 1
    assert client.get("/api/v1/governance-settings", headers=remediator_headers).status_code == 403
    audit_events = client.get("/api/v1/audit-events", headers=admin_headers).json()
    assert any(item["action"] == "GOVERNANCE_SETTING_UPDATED" and item["object_id"] == "sla" for item in audit_events)


def test_full_remediation_and_verification_workflow(client, admin_headers, remediator_headers, verifier_headers):
    finding = first_with_status(client, admin_headers, "pending_confirmation")
    users = client.get("/api/v1/users", headers=admin_headers).json()
    admin = next(item for item in users if "platform_admin" in item["roles"])
    remediator = next(item for item in users if "remediator" in item["roles"])
    verifier = next(item for item in users if "verifier" in item["roles"])
    assigned = client.patch(
        f"/api/v1/findings/{finding['id']}/assignment",
        headers=admin_headers,
        json={
            "owner_id": admin["id"],
            "assignee_id": remediator["id"],
            "verifier_id": verifier["id"],
            "due_at": (datetime.now(timezone.utc) + timedelta(days=7)).isoformat(),
            "reason": "测试流程分派",
            "version": finding["version"],
        },
    )
    assert assigned.status_code == 200, assigned.text
    confirmed = client.post(
        f"/api/v1/findings/{finding['id']}/transitions",
        headers=admin_headers,
        json={"action": "CONFIRM", "reason": "确认需要整改", "version": assigned.json()["version"]},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "pending_remediation"

    started = client.post(
        f"/api/v1/findings/{finding['id']}/transitions",
        headers=remediator_headers,
        json={"action": "START_REMEDIATION", "reason": "开始整改", "version": confirmed.json()["version"]},
    )
    assert started.status_code == 200, started.text
    assert started.json()["status"] == "in_remediation"

    submitted = client.post(
        f"/api/v1/findings/{finding['id']}/remediations",
        headers=remediator_headers,
        json={"description": "已完成修复并启用控制措施", "evidence": ["https://evidence.example/test"], "version": started.json()["version"]},
    )
    assert submitted.status_code == 200, submitted.text
    assert submitted.json()["status"] == "pending_verification"

    verified = client.post(
        f"/api/v1/findings/{finding['id']}/verifications",
        headers=verifier_headers,
        json={"result": "passed", "method": "manual", "comment": "复测通过", "version": submitted.json()["version"]},
    )
    assert verified.status_code == 200, verified.text
    assert verified.json()["status"] == "closed"
    events = client.get(f"/api/v1/findings/{finding['id']}/events", headers=admin_headers).json()
    assert {item["event_type"] for item in events} >= {"START_REMEDIATION", "REMEDIATION_SUBMITTED", "VERIFICATION_PASSED"}


def test_risk_acceptance_request_and_approval(client, admin_headers, remediator_headers):
    finding = first_with_status(client, remediator_headers, "pending_remediation")
    requested = client.post(
        f"/api/v1/findings/{finding['id']}/risk-acceptances",
        headers=remediator_headers,
        json={
            "reason": "业务窗口暂不可整改",
            "compensating_control": "限制网络访问并增加监控",
            "expires_at": (datetime.now(timezone.utc) + timedelta(days=20)).isoformat(),
            "version": finding["version"],
        },
    )
    assert requested.status_code == 200, requested.text
    assert requested.json()["status"] == "acceptance_requested"
    approved = client.post(
        f"/api/v1/findings/{finding['id']}/risk-acceptances/decision",
        headers=admin_headers,
        json={"approved": True, "comment": "同意临时接受", "version": requested.json()["version"]},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "risk_accepted"


def test_dashboard_reports_and_audit(client, admin_headers):
    summary = client.get("/api/v1/dashboard/summary", headers=admin_headers)
    report = client.get("/api/v1/reports/overview", headers=admin_headers)
    audit = client.get("/api/v1/audit-events", headers=admin_headers)
    assert summary.status_code == report.status_code == audit.status_code == 200
    assert "by_severity" in summary.json()
    assert "sla_compliance" in report.json()
    assert len(audit.json()) > 0
